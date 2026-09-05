"""XLSX parsing (spec section 10).

openpyxl is imported lazily so that the matching engine and the domain layer
remain importable in environments without it.

Cell values are converted to text using the *unformatted* value, except that
Excel dates are rendered as ISO so the date detector never has to guess.
Floats coming out of Excel are converted through ``repr`` and then ``Decimal``
rather than being handed to ``float(...)`` arithmetic anywhere downstream.
"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from typing import Any

from packages.ingestion.parsers.csv_parser import ParsedTable, _dedupe_headers

__all__ = ["list_sheets", "parse_xlsx"]


def _require_openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError(
            "openpyxl is required to read .xlsx files; install the 'recon-platform' extras"
        ) from exc
    return openpyxl


def list_sheets(data: bytes) -> list[str]:
    import io

    openpyxl = _require_openpyxl()
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        # Excel stores everything numeric as a double. Round-trip through the
        # shortest repr so 982.45 does not become 982.4499999999999.
        return str(Decimal(repr(value)).normalize())
    if isinstance(value, Decimal):
        return str(value)
    return str(value).strip()


def parse_xlsx(
    data: bytes,
    *,
    sheet_name: str | None = None,
    max_rows: int | None = None,
) -> ParsedTable:
    """Parse the first (or named) worksheet into a :class:`ParsedTable`."""
    import io

    openpyxl = _require_openpyxl()
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        rows_iter = worksheet.iter_rows(values_only=True)

        header_row: tuple[Any, ...] | None = None
        for candidate in rows_iter:
            if candidate and any(cell is not None and str(cell).strip() for cell in candidate):
                header_row = candidate
                break

        if header_row is None:
            return ParsedTable(
                columns=[],
                rows=[],
                encoding="xlsx",
                delimiter="",
                sheet_name=worksheet.title,
            )

        columns = _dedupe_headers([_cell_to_text(c) for c in header_row])
        rows: list[dict[str, str]] = []
        skipped: list[tuple[int, str]] = []
        total = 0

        for line_number, raw_row in enumerate(rows_iter, start=2):
            values = [_cell_to_text(c) for c in raw_row]
            if not any(v for v in values):
                continue
            total += 1
            if len(values) > len(columns):
                extra = [v for v in values[len(columns):] if v]
                if extra:
                    skipped.append(
                        (line_number, f"row has {len(values)} values for {len(columns)} columns")
                    )
                    continue
                values = values[: len(columns)]
            while len(values) < len(columns):
                values.append("")
            rows.append(dict(zip(columns, values, strict=True)))
            if max_rows is not None and len(rows) >= max_rows:
                break

        return ParsedTable(
            columns=columns,
            rows=rows,
            encoding="xlsx",
            delimiter="",
            sheet_name=worksheet.title,
            total_rows=total,
            skipped_rows=skipped,
        )
    finally:
        workbook.close()
